"""
Shadow backtest for the coil/session strategy on a single IST date. NO orders,
NO live-order code path -- pure simulation on 5-minute bars.

    python coil_backtest.py                       # yesterday (IST)
    python coil_backtest.py --date 2026-06-03 -v
    python coil_backtest.py --risk 50 --breakeven 30 --daily-limit 150 --units 30000

Per session: run the real coil selection, then replay the selected pair's session
bars through the bias/DMI state machine (`coil_strategy`) with a simulated risk
layer (stop-market, live-... approximated by bar extremes, breakeven latch,
daily flatten+halt). See SPEC_COIL.md §6 and §9 for the rules and the
backtest approximations.

USD amounts are placeholders unless you pass --risk/--breakeven/--daily-limit.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import logging
from datetime import date, datetime, timedelta
from typing import List, Optional

from coil_config import load_coil_config
from coil_selection import coil_select
from coil_strategy import CoilBase, entry_dir, reverse_signal
from config import IBKRConnection, StrategyConfig
from ibkr_client import IBKRClient
from pivots import compute_pivots
from quality_supertrend import DMI
from selection import SelectionError
from sessions import IST, Session, first_candle_close, session_window
from time_utils import NY, to_ny

# Reuse the evaluator's fetch+select helpers (already validated against the gateway).
from coil_select_eval import _bar_open_ny, _cpr_from_prior_fx_day, _metrics_for_session

logger = logging.getLogger("coil_bt")

WARMUP_DAYS = 3          # DMI warm-up before the session start


def _write_excel(path: str, blotter: list[dict], day_state: dict, d: date, args,
                 bars_all: list[dict]) -> None:
    """Write the blotter to an .xlsx: Trades + Bars (OHLC+DMI) + Summary sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    def pr(v, sym):
        """Round a price: JPY pairs -> 3 dp, others -> 5 dp."""
        if v is None:
            return ""
        return round(v, 3 if str(sym).upper().endswith("JPY") else 5)

    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"
    headers = ["Session", "Time (IST)", "Symbol", "Action", "Side",
               "Price", "PnL (USD)", "Running PnL", "Base", "Reason"]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center")

    running = 0.0
    per_session: dict[str, float] = {}
    for r in blotter:
        if "note" in r:
            ws.append([r["session"], "", "", "—", "", "", "", "", "", r["note"]])
            continue
        running += r["pnl"]
        per_session[r["session"]] = per_session.get(r["session"], 0.0) + r["pnl"]
        side = "LONG" if r["side"] == "L" else "SHORT"
        ws.append([r["session"], r["t"], r["sym"], r["action"], side,
                   pr(r["px"], r["sym"]), round(r["pnl"], 2), round(running, 2),
                   r["base"], r["reason"]])
        if r["pnl"]:
            tint = "C6EFCE" if r["pnl"] > 0 else "FFC7CE"
            ws.cell(row=ws.max_row, column=7).fill = PatternFill("solid", fgColor=tint)

    widths = [16, 11, 9, 18, 7, 12, 11, 12, 7, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    # --- Bars sheet: every 5m session bar with OHLC + DMI for verification ---
    bs = wb.create_sheet("Bars")
    bhead = ["Session", "Time (IST)", "Symbol", "Open", "High", "Low", "Close",
             "DI+", "DI-", "ADX", "Base", "Base Px", "Bias", "Position", "Action(s)"]
    bs.append(bhead)
    for cell in bs[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center")
    cur_sess = None
    pos_txt = {1: "LONG", -1: "SHORT", 0: ""}
    for r in bars_all:
        if cur_sess is not None and r["session"] != cur_sess:
            bs.append([])                      # blank row between sessions
        cur_sess = r["session"]
        rd = lambda v, n: ("" if v is None else round(v, n))
        sym = r["sym"]
        bs.append([r["session"], r["t"], sym,
                   pr(r["o"], sym), pr(r["h"], sym), pr(r["l"], sym), pr(r["c"], sym),
                   rd(r["plus"], 4), rd(r["minus"], 4), rd(r["adx"], 4),
                   r["base"] or "", pr(r["base_px"], sym), r["bias"],
                   pos_txt.get(r["pos"], ""), r["acts"]])
        if r["acts"]:
            for col in range(1, len(bhead) + 1):
                bs.cell(row=bs.max_row, column=col).fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate([16, 11, 9, 11, 11, 11, 11, 8, 8, 8, 7, 11, 7, 10, 22], 1):
        bs.column_dimensions[chr(64 + i)].width = w
    bs.freeze_panes = "A2"

    sm = wb.create_sheet("Summary")
    sm.append([f"Coil backtest {d}"])
    sm["A1"].font = Font(bold=True, size=13)
    sm.append(["units", args.units]); sm.append(["risk/trade ($)", args.risk])
    sm.append(["breakeven ($)", args.breakeven]); sm.append(["daily limit ($)", args.daily_limit])
    sm.append([])
    sm.append(["Session", "Net PnL ($)"]); sm["A7"].font = sm["B7"].font = Font(bold=True)
    for name, net in per_session.items():
        sm.append([name, round(net, 2)])
    sm.append([])
    sm.append(["DAY REALIZED", round(day_state["realized"], 2)])
    sm.append(["Halted on daily limit", "YES" if day_state["halted"] else "NO"])
    for col in ("A", "B"):
        sm.column_dimensions[col].width = 22
    wb.save(path)
CONV = {"JPY": "USDJPY", "CHF": "USDCHF", "CAD": "USDCAD"}


async def _usd_per_price_unit(ibkr: IBKRClient, symbol: str, at_ny: datetime, units: int) -> float:
    """
    USD value of a 1.0 price move on `units` units of `symbol`.
    = units x (quote->USD). Quote=USD -> 1; JPY/CHF/CAD -> 1 / USD{ccy} rate.
    """
    quote = symbol[3:].upper()
    if quote == "USD":
        return float(units)
    conv = CONV.get(quote)
    if conv is None:
        raise SelectionError(f"no USD conversion known for quote {quote}")
    bars = await ibkr.fetch_5min_bars(conv, end_ny=at_ny, duration_str="1 D")
    if not bars:
        raise SelectionError(f"no {conv} bars to convert {quote}->USD")
    rate = float(bars[-1].close)            # USD per 1 {ccy} = 1/rate
    return float(units) / rate


async def _run_session(ibkr: IBKRClient, session: Session, d: date, params, units, risk,
                       breakeven, day_state, bars_all, commission) -> list[dict]:
    """Select + replay one session for the selected pair. Appends to blotter rows."""
    rows: list[dict] = []
    # --- 1. selection (real) ---
    metrics = await _metrics_for_session(ibkr, session, d, params, use_cfd=False)
    try:
        winner = coil_select(metrics, adx_filter_max=params.adx_filter_max).symbol
    except SelectionError as e:
        rows.append({"session": session.name, "note": f"no selection ({e})"})
        return rows

    start_ist, end_ist = session_window(session, d)
    start_ny, end_ny = start_ist.astimezone(NY), end_ist.astimezone(NY)
    anchor_ny = first_candle_close(session, d).astimezone(NY)
    noentry_cutoff_ny = end_ny - timedelta(minutes=params.no_entry_last_minutes)

    # --- 2. data: pivots (prior FX day) + warm-up + session bars ---
    bars = await ibkr.fetch_5min_bars_range(
        winner, start_ny=start_ny - timedelta(days=WARMUP_DAYS), end_ny=end_ny, use_cfd=False)
    cpr = _cpr_from_prior_fx_day(bars, anchor_ny)
    piv = compute_pivots(cpr.high, cpr.low, cpr.close)
    warm = [b for b in bars if _bar_open_ny(b) < start_ny]
    sess = [b for b in bars if start_ny <= _bar_open_ny(b) < end_ny]
    if not sess:
        rows.append({"session": session.name, "note": f"{winner}: no session bars"})
        return rows

    # BID/ASK session bars -> REAL fills: buy at the ASK, sell at the BID (models
    # the spread crossed on every market order).
    bid_bars = await ibkr.fetch_5min_bars_range(
        winner, start_ny=start_ny, end_ny=end_ny, use_cfd=False, what="BID")
    ask_bars = await ibkr.fetch_5min_bars_range(
        winner, start_ny=start_ny, end_ny=end_ny, use_cfd=False, what="ASK")
    bid_close = {_bar_open_ny(b).isoformat(): float(b.close) for b in bid_bars}
    ask_close = {_bar_open_ny(b).isoformat(): float(b.close) for b in ask_bars}

    upu = await _usd_per_price_unit(ibkr, winner, start_ny, units)   # USD per 1.0 price move
    stop_dist = risk / upu                                            # price units
    pnl = lambda side, entry, px: side * (px - entry) * upu

    # --- 3. replay ---
    dmi = DMI(14, 14)
    for b in warm:
        dmi.update(float(b.high), float(b.low), float(b.close))

    base = CoilBase()
    pos = 0
    entry = stop = None
    be_armed = False

    def realize(px, action, ts, reason):
        nonlocal pos, entry, stop, be_armed
        # gross PnL minus round-trip commission (open + close = 2 orders).
        p = pnl(pos, entry, px) - 2.0 * commission
        day_state["realized"] += p
        rows.append({"session": session.name, "sym": winner, "t": ts, "action": action,
                     "side": "L" if pos > 0 else "S", "px": px, "pnl": p,
                     "base": f"{base.name}", "reason": reason})
        pos = 0; entry = None; stop = None; be_armed = False
        return p

    n = len(sess)
    for i, b in enumerate(sess):
        o, h, l, c = float(b.open), float(b.high), float(b.low), float(b.close)
        open_ny = _bar_open_ny(b)
        ts = open_ny.astimezone(IST).strftime("%H:%M")
        ask_px = ask_close.get(open_ny.isoformat(), c)   # buy fills here
        bid_px = bid_close.get(open_ny.isoformat(), c)   # sell fills here
        plus, minus, adx = dmi.update(h, l, c)
        is_first = base.update(h, l, c, piv)
        is_end = (i == n - 1)
        # Final-hour lockout: no NEW entries / reversal re-entries (exits still fire).
        no_new = is_end or open_ny >= noentry_cutoff_ny

        def _log_bar(row_start: int):
            acts = "+".join(r["action"] for r in rows[row_start:] if "action" in r)
            if is_first:
                acts = ("SEED " + acts).strip()
            bars_all.append({
                "session": session.name, "sym": winner, "t": ts,
                "o": o, "h": h, "l": l, "c": c,
                "plus": plus, "minus": minus, "adx": adx,
                "base": base.name, "base_px": base.level, "bias": base.bias(c),
                "pos": pos, "acts": acts,
            })

        row_start = len(rows)
        if plus is None:
            _log_bar(row_start)
            continue

        # (1) intrabar protective exit, checked first against the candle wick:
        #     the per-trade stop AND the daily-cap level (realized + UNREALIZED),
        #     whichever the adverse wick reaches first. Exit at that exact price.
        if pos != 0 and stop is not None:
            remaining = day_state["limit"] + day_state["realized"]   # $ left before the daily cap
            daily_stop = entry - remaining / upu if pos > 0 else entry + remaining / upu
            if pos > 0:
                binding = max(stop, daily_stop); is_daily = daily_stop >= stop
            else:
                binding = min(stop, daily_stop); is_daily = daily_stop <= stop
            adverse = l if pos > 0 else h
            if (pos > 0 and adverse <= binding) or (pos < 0 and adverse >= binding):
                realize(binding, "EXIT_DAILY_LIMIT" if is_daily else "STOP_OUT", ts,
                        "daily_limit" if is_daily else "stop")
                if is_daily:
                    day_state["halted"] = True
                    _log_bar(row_start)
                    break
        # (2) breakeven latch on favourable extreme
        if pos != 0 and not be_armed:
            fav = h if pos > 0 else l
            if pnl(pos, entry, fav) >= breakeven:
                stop = entry; be_armed = True
                rows.append({"session": session.name, "sym": winner, "t": ts,
                             "action": "BE_MOVE", "side": "L" if pos > 0 else "S",
                             "px": entry, "pnl": 0.0, "base": f"{base.name}", "reason": "breakeven"})
        # (3) signal at close: reverse / entry
        bias = base.bias(c)
        if pos != 0 and reverse_signal(pos, bias, plus, minus):
            old = pos
            realize(bid_px if old > 0 else ask_px, "EXIT_SIGNAL", ts, "opp_signal")  # close long@bid / short@ask
            if not no_new:
                pos = -old
                entry = ask_px if pos > 0 else bid_px                # open long@ask / short@bid
                stop = entry - stop_dist if pos > 0 else entry + stop_dist
                rows.append({"session": session.name, "sym": winner, "t": ts, "action": "REVERSE",
                             "side": "L" if pos > 0 else "S", "px": entry, "pnl": 0.0,
                             "base": f"{base.name}", "reason": "stop_reverse"})
        if pos == 0 and not is_first and not no_new:
            d_ = entry_dir(bias, plus, minus)
            if d_:
                pos = d_
                entry = ask_px if pos > 0 else bid_px               # open long@ask / short@bid
                stop = entry - stop_dist if pos > 0 else entry + stop_dist
                rows.append({"session": session.name, "sym": winner, "t": ts, "action": "ENTRY",
                             "side": "L" if pos > 0 else "S", "px": entry, "pnl": 0.0,
                             "base": f"{base.name}", "reason": "gates_align"})
        # (4) session-end force-flat
        if is_end and pos != 0:
            realize(bid_px if pos > 0 else ask_px, "EXIT_EOD", ts, "session_end")  # close long@bid / short@ask

        _log_bar(row_start)

        # (5) daily limit: flatten + halt
        if day_state["realized"] <= -day_state["limit"]:
            if pos != 0:
                realize(c, "EXIT_DAILY_LIMIT", ts, "daily_limit")
            day_state["halted"] = True
            break
    return rows


async def run(args) -> None:
    sessions, params, _feed = load_coil_config(args.config)
    d = (datetime.strptime(args.date, "%Y-%m-%d").date() if args.date
         else (datetime.now(IST) - timedelta(days=1)).date())
    if getattr(args, "offline", False):
        from bar_cache import CachedBars
        ibkr = CachedBars()                      # reads bars from local cache, no IBKR
    else:
        cfg = StrategyConfig(ibkr=IBKRConnection(
            host=args.host, port=args.port, client_id=args.client_id, read_only=True))
        ibkr = IBKRClient(cfg)
    await ibkr.connect()
    day_state = {"realized": 0.0, "limit": args.daily_limit, "halted": False}
    blotter: list[dict] = []
    bars_all: list[dict] = []
    try:
        for session in sessions:
            if day_state["halted"]:
                blotter.append({"session": session.name, "note": "halted (daily limit)"})
                continue
            blotter += await _run_session(ibkr, session, d, params, args.units,
                                          args.risk, args.breakeven, day_state, bars_all,
                                          args.commission)
    finally:
        await ibkr.disconnect()

    _emit(d, args, day_state, blotter, bars_all, save_cache=True)


def _emit(d, args, day_state, blotter, bars_all, save_cache: bool) -> None:
    """Print the blotter, write the Excel, and (optionally) cache the raw data
    so the workbook can be re-rendered offline (no Gateway) via --render-only."""
    print(f"\n=== Coil backtest {d}  units={args.units}  risk=${args.risk:g}  "
          f"BE=${args.breakeven:g}  daily_limit=${args.daily_limit:g} ===")
    print("(placeholder USD params unless overridden; bar-based stop/BE sim - see SPEC_COIL.md §9)\n")
    cur = None
    for r in blotter:
        if r["session"] != cur:
            cur = r["session"]; print(f"--- {cur} ---")
        if "note" in r:
            print(f"    {r['note']}"); continue
        print(f"    {r['t']}  {r['action']:<16} {r['side']} {r['sym']:<7} @ {r['px']:.5f}"
              f"  base={r['base']:<3}  pnl={r['pnl']:+8.2f}  ({r['reason']})")
    print(f"\n  DAY REALIZED PnL: ${day_state['realized']:+.2f}"
          + ("   [HALTED on daily limit]" if day_state["halted"] else ""))

    if save_cache:
        cache = f"coil_bt_cache_{d}.json"
        with open(cache, "w") as f:
            json.dump({"date": str(d),
                       "args": {"units": args.units, "risk": args.risk,
                                "breakeven": args.breakeven, "daily_limit": args.daily_limit},
                       "day_state": day_state, "blotter": blotter, "bars": bars_all}, f)
        print(f"  Cache written: {cache}  (re-render offline with --render-only)")

    xlsx = args.excel or f"coil_trades_{d}.xlsx"
    _write_excel(xlsx, blotter, day_state, d, args, bars_all)
    print(f"  Excel written: {xlsx}  (sheets: Trades, Bars, Summary)")


def main() -> int:
    p = argparse.ArgumentParser(description="Coil/session strategy shadow backtest (no orders)")
    p.add_argument("--date", help="IST date YYYY-MM-DD (default: yesterday IST)")
    p.add_argument("--config", help="JSON config (sessions + params)")
    p.add_argument("--offline", action="store_true", help="Use local bar cache (no IBKR)")
    p.add_argument("--excel", help="Output .xlsx path (default: coil_trades_<date>.xlsx)")
    p.add_argument("--render-only", help="Rebuild the Excel from a cache JSON (no Gateway)")
    p.add_argument("--units", type=int, default=30000, help="Position size (base-ccy units)")
    p.add_argument("--risk", type=float, default=50.0, help="USD risked per trade")
    p.add_argument("--breakeven", type=float, default=50.0, help="USD profit that arms breakeven")
    p.add_argument("--daily-limit", type=float, default=150.0, help="USD daily loss cap")
    p.add_argument("--commission", type=float, default=2.0, help="USD commission per order ($/fill)")
    p.add_argument("--host", default="127.0.0.1")
    p.add_argument("--port", type=int, default=4001)
    p.add_argument("--client-id", type=int, default=71)
    p.add_argument("--verbose", "-v", action="store_true")
    args = p.parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.WARNING,
                        format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
    if args.render_only:
        with open(args.render_only) as f:
            data = json.load(f)
        a = data["args"]
        args.units, args.risk = a["units"], a["risk"]
        args.breakeven, args.daily_limit = a["breakeven"], a["daily_limit"]
        _emit(date.fromisoformat(data["date"]), args, data["day_state"],
              data["blotter"], data["bars"], save_cache=False)
        return 0
    asyncio.run(run(args))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
