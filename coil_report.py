"""
Consolidate per-day backtest caches (coil_bt_cache_<date>.json) into ONE Excel
report: every trade, daily P&L, and full summary stats. Offline (no Gateway).

    python coil_report.py 2026-04-09 2026-06-05 [out.xlsx]
"""
from __future__ import annotations
import glob, json, sys
from collections import defaultdict

OPEN_ACTIONS = {"ENTRY", "REVERSE"}
CLOSE_ACTIONS = {"EXIT_SIGNAL", "STOP_OUT", "EXIT_EOD", "EXIT_DAILY_LIMIT"}


def collect(start: str, end: str):
    trades, daily = [], []
    for cf in sorted(glob.glob("coil_bt_cache_*.json")):
        d = json.load(open(cf))
        day = d["date"]
        if not (start <= day <= end):
            continue
        blot = d["blotter"]
        # pair each entry with its exit -> one row per completed trade
        op = None
        n_trades = 0
        for r in blot:
            a = r.get("action")
            if a in OPEN_ACTIONS:
                op = r
            elif a in CLOSE_ACTIONS and op is not None:
                trades.append({
                    "date": day, "session": r["session"], "sym": r.get("sym", ""),
                    "side": "LONG" if op.get("side") == "L" else "SHORT",
                    "entry_t": op.get("t"), "exit_t": r.get("t"),
                    "entry": op.get("px"), "exit": r.get("px"),
                    "pnl": round(r.get("pnl", 0.0), 2), "exit_reason": r.get("reason", ""),
                })
                n_trades += 1
                op = None
        daily.append({"date": day, "realized": round(d["day_state"]["realized"], 2),
                      "halted": d["day_state"]["halted"], "trades": n_trades})
    return trades, sorted(daily, key=lambda x: x["date"])


def write(path, trades, daily, start, end):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    hf = PatternFill("solid", fgColor="1F4E78")
    def hdr(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = hf; c.alignment = Alignment(horizontal="center")

    wb = Workbook()
    # --- Summary ---
    sm = wb.active; sm.title = "Summary"
    pnls = [t["pnl"] for t in trades]
    wins = [p for p in pnls if p > 0]; losses = [p for p in pnls if p < 0]
    total = sum(d["realized"] for d in daily)
    up_days = [d for d in daily if d["realized"] > 0]; down_days = [d for d in daily if d["realized"] < 0]
    # equity curve / max drawdown
    eq = 0.0; peak = 0.0; mdd = 0.0
    for d in daily:
        eq += d["realized"]; peak = max(peak, eq); mdd = min(mdd, eq - peak)
    rows = [
        ("Period", f"{start} -> {end}"),
        ("Trading days", len(daily)),
        ("Total realized P&L ($)", round(total, 2)),
        ("", ""),
        ("Total trades", len(trades)),
        ("Wins / Losses", f"{len(wins)} / {len(losses)}"),
        ("Win rate", f"{(len(wins)/len(trades)*100):.1f}%" if trades else "-"),
        ("Avg win ($)", round(sum(wins)/len(wins), 2) if wins else 0),
        ("Avg loss ($)", round(sum(losses)/len(losses), 2) if losses else 0),
        ("Largest win ($)", round(max(pnls), 2) if pnls else 0),
        ("Largest loss ($)", round(min(pnls), 2) if pnls else 0),
        ("Profit factor", round(sum(wins)/abs(sum(losses)), 2) if losses else "-"),
        ("", ""),
        ("Up days / Down days", f"{len(up_days)} / {len(down_days)}"),
        ("Best day ($)", max((d["realized"] for d in daily), default=0)),
        ("Worst day ($)", min((d["realized"] for d in daily), default=0)),
        ("Days hit daily cap", sum(1 for d in daily if d["halted"])),
        ("Max drawdown ($)", round(mdd, 2)),
    ]
    sm.append(["COIL STRATEGY - 2 MONTH BACKTEST"]); sm["A1"].font = Font(bold=True, size=14)
    sm.append([])
    for k, v in rows:
        sm.append([k, v])
    # per-session and per-month
    sm.append([]); sm.append(["By session", "trades", "P&L ($)", "win%"]);
    bys = defaultdict(lambda: [0, 0.0, 0])
    for t in trades:
        s = bys[t["session"]]; s[0] += 1; s[1] += t["pnl"]; s[2] += 1 if t["pnl"] > 0 else 0
    for name, (n, p, w) in bys.items():
        sm.append([name, n, round(p, 2), f"{(w/n*100):.0f}%" if n else "-"])
    sm.append([]); sm.append(["By month", "P&L ($)", "trades"])
    bym = defaultdict(lambda: [0.0, 0])
    for d in daily:
        bym[d["date"][:7]][0] += d["realized"]
    for t in trades:
        bym[t["date"][:7]][1] += 1
    for m in sorted(bym):
        sm.append([m, round(bym[m][0], 2), bym[m][1]])
    for col in ("A", "B", "C", "D"):
        sm.column_dimensions[col].width = 22

    # --- Daily ---
    dd = wb.create_sheet("Daily"); hdr(dd, ["Date", "Realized P&L ($)", "Trades", "Halted"])
    run = 0.0
    for d in daily:
        run += d["realized"]
        dd.append([d["date"], d["realized"], d["trades"], "YES" if d["halted"] else "", round(run, 2)])
    dd.cell(row=1, column=5, value="Cumulative ($)").font = Font(bold=True, color="FFFFFF")
    dd.cell(row=1, column=5).fill = hf
    for i, w in enumerate([12, 16, 8, 8, 14], 1):
        dd.column_dimensions[chr(64+i)].width = w
    dd.freeze_panes = "A2"

    # --- Trades ---
    tr = wb.create_sheet("Trades")
    hdr(tr, ["Date", "Session", "Symbol", "Side", "Entry t", "Exit t", "Entry", "Exit", "P&L ($)", "Exit reason"])
    for t in trades:
        dp = 3 if str(t["sym"]).endswith("JPY") else 5
        tr.append([t["date"], t["session"], t["sym"], t["side"], t["entry_t"], t["exit_t"],
                   round(t["entry"], dp) if t["entry"] else "", round(t["exit"], dp) if t["exit"] else "",
                   t["pnl"], t["exit_reason"]])
        if t["pnl"]:
            tr.cell(row=tr.max_row, column=9).fill = PatternFill(
                "solid", fgColor="C6EFCE" if t["pnl"] > 0 else "FFC7CE")
    for i, w in enumerate([12, 15, 9, 7, 8, 8, 11, 11, 10, 14], 1):
        tr.column_dimensions[chr(64+i)].width = w
    tr.freeze_panes = "A2"

    # --- one sub-sheet per session: stats block + that session's trades ---
    by_sess = defaultdict(list)
    for t in trades:
        by_sess[t["session"]].append(t)
    for sname, st in by_sess.items():
        ws = wb.create_sheet(sname[:31])
        sp = [t["pnl"] for t in st]
        sw = [p for p in sp if p > 0]; sl = [p for p in sp if p < 0]
        byday = defaultdict(float)
        for t in st:
            byday[t["date"]] += t["pnl"]
        ws.append([f"{sname}  —  summary"]); ws["A1"].font = Font(bold=True, size=13)
        for k, v in [
            ("Total P&L ($)", round(sum(sp), 2)),
            ("Trades", len(st)),
            ("Wins / Losses", f"{len(sw)} / {len(sl)}"),
            ("Win rate", f"{len(sw)/len(st)*100:.1f}%" if st else "-"),
            ("Avg win / loss ($)", f"{round(sum(sw)/len(sw),2) if sw else 0} / "
                                   f"{round(sum(sl)/len(sl),2) if sl else 0}"),
            ("Profit factor", round(sum(sw)/abs(sum(sl)), 2) if sl else "-"),
            ("Days traded", len(byday)),
            ("Best / worst day ($)", f"{round(max(byday.values()),2)} / "
                                     f"{round(min(byday.values()),2)}" if byday else "-"),
        ]:
            ws.append([k, v])
        ws.append([])
        ws.append(["Date", "Symbol", "Side", "Entry t", "Exit t", "Entry", "Exit", "P&L ($)", "Exit reason"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = hf; c.alignment = Alignment(horizontal="center")
        hdr_row = ws.max_row
        for t in st:
            dp = 3 if str(t["sym"]).endswith("JPY") else 5
            ws.append([t["date"], t["sym"], t["side"], t["entry_t"], t["exit_t"],
                       round(t["entry"], dp) if t["entry"] else "", round(t["exit"], dp) if t["exit"] else "",
                       t["pnl"], t["exit_reason"]])
            if t["pnl"]:
                ws.cell(row=ws.max_row, column=8).fill = PatternFill(
                    "solid", fgColor="C6EFCE" if t["pnl"] > 0 else "FFC7CE")
        for i, w in enumerate([12, 9, 7, 8, 8, 11, 11, 10, 14], 1):
            ws.column_dimensions[chr(64+i)].width = w
        ws.freeze_panes = f"A{hdr_row+1}"

    wb.save(path)
    return len(trades), round(total, 2)


if __name__ == "__main__":
    start = sys.argv[1] if len(sys.argv) > 1 else "2026-04-09"
    end = sys.argv[2] if len(sys.argv) > 2 else "2026-06-05"
    out = sys.argv[3] if len(sys.argv) > 3 else "coil_2month_report.xlsx"
    trades, daily = collect(start, end)
    n, total = write(out, trades, daily, start, end)
    print(f"wrote {out}: {len(daily)} days, {n} trades, total P&L ${total:+.2f}")
